"""
ETL script: extracts data from GoldBod_Financial_Analysis.xlsx
into clean CSV files for Power BI.

Run with the venv active:
    python etl.py
"""

import openpyxl
import csv
import os

SOURCE_FILE = "GoldBod_Financial_Analysis.xlsx"
OUTPUT_FOLDER = "output"

STANDARD_SHEETS = ["Overview", "Income Statement", "Balance Sheet", "Key Ratios"]


def extract_standard_sheet(ws):
    """Reads label, FY2025, FY2024, change/reading, and note columns."""
    rows = []

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        label = row[0].value
        fy2025 = row[1].value if len(row) > 1 else None

        if label is None:
            continue
        if not isinstance(fy2025, (int, float)):
            continue

        fy2024 = row[2].value if len(row) > 2 else None
        change_or_reading = row[3].value if len(row) > 3 else None
        note = row[4].value if len(row) > 4 else None

        rows.append({
            "metric": str(label).strip(),
            "fy2025": fy2025,
            "fy2024": fy2024 if isinstance(fy2024, (int, float)) else None,
            "yoy_change_or_reading": change_or_reading,
            "note": note,
        })

    return rows


def extract_imf_timeline(ws):
    """Reads the Date, Event, Figure cited, Source table."""
    rows = []
    header_found = False

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        a = row[0].value
        b = row[1].value
        c = row[2].value
        d = row[3].value

        if a == "Date" and b == "Event":
            header_found = True
            continue

        if not header_found:
            continue
        if a is None or b is None:
            continue
        if c is None and d is None:
            break

        rows.append({
            "date": a,
            "event": b,
            "figure_cited": c,
            "source": d,
        })

    return rows


def write_csv(filename, rows):
    if not rows:
        print(f"  no rows found, skipping {filename}")
        return

    path = os.path.join(OUTPUT_FOLDER, filename)

    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)

    print(f"  wrote {len(rows)} rows to {path}")


def main():
    if not os.path.exists(SOURCE_FILE):
        raise FileNotFoundError(f"Can't find {SOURCE_FILE}. Put it in this folder.")

    os.makedirs(OUTPUT_FOLDER, exist_ok=True)

    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)

    for sheet_name in STANDARD_SHEETS:
        print(f"Extracting: {sheet_name}")
        ws = wb[sheet_name]
        rows = extract_standard_sheet(ws)
        out_name = sheet_name.lower().replace(" ", "_") + ".csv"
        write_csv(out_name, rows)

    print("Extracting: IMF Controversy Context")
    ws = wb["IMF Controversy Context"]
    imf_rows = extract_imf_timeline(ws)
    write_csv("imf_timeline.csv", imf_rows)

    print("\nDone. Check the output folder.")


if __name__ == "__main__":
    main()